"""Safe, data-only Master workbook parsing and validation.

This module deliberately has no database dependency. Parsing and validation are
completed before a caller opens a transaction or changes business data.
"""

import csv
import hashlib
import hmac
import io
import json
import re
import time
import uuid
from datetime import date, datetime

import openpyxl

MAX_BYTES = 10 * 1024 * 1024
ALLOWED_SUFFIXES = {".csv", ".xlsx", ".xlsm"}
PREVIEW_TTL_SECONDS = 15 * 60

# Aliases are deliberately keyed by table/master. Generic parsing remains
# independent of spreadsheet-specific workbook layouts.
ALIASES = {
    "companies": {"id": "company_id", "company_id": "company_id", "company_name": "company_name", "company": "company_name"},
    "company_entities": {"id": "entity_id", "entity_id": "entity_id", "entity_code": "entity_code"},
    "employees": {"employee_id": "emp_id", "employee_code": "emp_code", "name": "guest_name", "employee_name": "guest_name", "email": "guest_email", "mobile": "guest_mobile"},
    "individuals": {"id": "individual_id", "individual_id": "individual_id", "name": "guest_name", "mobile": "guest_contact", "email": "guest_email"},
    "contacts": {"id": "contact_id", "admin_id": "contact_id", "name": "contact_name", "email_id": "email", "phone": "mobile"},
    "contracts": {"id": "contract_id", "contract_number": "contract_no", "contract_id": "contract_id"},
    "ratecards": {"id": "rate_card_id", "rate_card_id": "rate_card_id", "vendor": "vendor_id", "company": "company_id"},
    "vendors": {"id": "vendor_id", "vendor_id": "vendor_id", "vendor": "vendor_name", "name": "vendor_name"},
    "vehicles": {"id": "vehicle_id", "vehicle_id": "vehicle_id", "registration_number": "reg_number", "registration_no": "reg_number"},
    "drivers": {"id": "driver_id", "driver_id": "driver_id", "name": "driver_name", "phone": "mobile"},
    "leads": {"id": "lead_id", "lead_id": "lead_id", "company_name": "company"},
    "settings": {"name": "setting_name", "key": "setting_name"},
}


def _name(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _text(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip() or None


def source_hash(content):
    return hashlib.sha256(content).hexdigest()


def _canonical(value):
    return json.dumps(value, sort_keys=True, separators=(",", ":"), default=str).encode()


def sign_preview(plan, secret, ttl=PREVIEW_TTL_SECONDS):
    """Create an opaque HMAC-signed preview token."""
    now = int(time.time())
    body = dict(plan, preview_id=plan.get("preview_id") or uuid.uuid4().hex,
                issued_at=now, expires_at=now + ttl)
    encoded = __import__("base64").urlsafe_b64encode(_canonical(body)).decode().rstrip("=")
    signature = hmac.new(secret.encode(), encoded.encode(), hashlib.sha256).hexdigest()
    return encoded + "." + signature


def verify_preview(token, secret, now=None):
    try:
        encoded, signature = token.rsplit(".", 1)
        expected = hmac.new(secret.encode(), encoded.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(signature, expected):
            return None, "SIGNATURE_INVALID"
        padded = encoded + "=" * (-len(encoded) % 4)
        body = json.loads(__import__("base64").urlsafe_b64decode(padded))
        if int(now or time.time()) > int(body["expires_at"]):
            return None, "PREVIEW_EXPIRED"
        return body, None
    except (ValueError, KeyError, TypeError, json.JSONDecodeError):
        return None, "SIGNATURE_INVALID"


def map_headers(headers, cfg):
    fields = {field["name"] for field in cfg.get("fields", [])} | {cfg["pk"], "tenant_id"}
    aliases = dict(ALIASES.get(cfg.get("table", ""), {}))
    aliases.update({_name(field): field for field in fields})
    for field in cfg.get("fields", []):
        for alias in field.get("aliases", []):
            aliases[_name(alias)] = field["name"]
    mapped, mapping, errors = [], {}, []
    for source in headers:
        target = aliases.get(_name(source))
        if not target:
            mapped.append(source)
            errors.append({"row": 1, "source_field": source, "code": "UNKNOWN_COLUMN",
                           "column": source, "message": "Unknown column; map it explicitly or remove it."})
            continue
        if target in mapping.values():
            errors.append({"row": 1, "source_field": source, "canonical_field": target,
                           "code": "AMBIGUOUS_MAPPING", "column": source,
                           "message": f"Multiple source columns map to {target}."})
            continue
        mapping[source] = target
        mapped.append(target)
    return mapped, mapping, errors


def read_rows(content, filename, sheet=None):
    """Return (headers, rows, errors), without evaluating formulas."""
    errors = []
    suffix = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if len(content) > MAX_BYTES:
        return [], [], [{"row": 0, "column": "file", "message": "File exceeds 10 MB."}]
    if suffix not in ALLOWED_SUFFIXES:
        return [], [], [{"row": 0, "column": "file", "message": "Only CSV, XLSX, and XLSM files are supported."}]
    try:
        if suffix == ".csv":
            values = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        else:
            book = openpyxl.load_workbook(io.BytesIO(content), read_only=True,
                                          data_only=True, keep_links=False)
            selected = sheet or (book.sheetnames[0] if book.sheetnames else None)
            if not selected or selected not in book.sheetnames:
                return [], [], [{"row": 0, "column": "sheet", "message": "Worksheet was not found."}]
            values = [list(row) for row in book[selected].iter_rows(values_only=True)]
    except Exception as exc:
        return [], [], [{"row": 0, "column": "file", "message": f"Could not read file: {exc}"}]
    if not values:
        return [], [], [{"row": 1, "column": "header", "message": "File contains no rows."}]
    headers = [_name(value) for value in values[0]]
    if not any(headers):
        errors.append({"row": 1, "column": "header", "message": "Header row is empty."})
    if len(values) == 1:
        errors.append({"row": 1, "column": "file", "code": "NO_DATA_ROWS", "message": "File contains a header but no data rows."})
    if len(headers) != len(set(headers)):
        errors.append({"row": 1, "column": "header", "message": "Header names must be unique."})
    rows = []
    for number, values_row in enumerate(values[1:], start=2):
        values_row = list(values_row) + [None] * (len(headers) - len(values_row))
        row = {header: _text(value) for header, value in zip(headers, values_row)}
        if any(value is not None for value in row.values()):
            row["_row"] = number
            rows.append(row)
    return headers, rows, errors


def validate_rows(headers, rows, cfg, tenant_id=None):
    """Validate against the existing config fields and primary key."""
    errors = []
    fields = {field["name"]: field for field in cfg.get("fields", [])}
    expected = set(fields) | {cfg["pk"], "tenant_id"}
    missing = sorted(name for name in expected if name not in headers and name == cfg["pk"])
    for name in missing:
        errors.append({"row": 1, "column": name, "code": "REQUIRED_FIELD_MISSING", "message": "Required key column is missing."})
    seen = set()
    for row in rows:
        number = row["_row"]
        key = row.get(cfg["pk"])
        if not key:
            errors.append({"row": number, "column": cfg["pk"], "code": "REQUIRED_FIELD_MISSING", "message": "Primary key is required."})
        elif key in seen:
            errors.append({"row": number, "column": cfg["pk"], "code": "DUPLICATE_KEY", "message": "Duplicate primary key in file."})
        seen.add(key)
        if row.get("tenant_id") and tenant_id and str(row["tenant_id"]).strip() != str(tenant_id).strip():
            errors.append({"row": number, "column": "tenant_id", "value": row["tenant_id"],
                           "code": "TENANT_MISMATCH", "message": "Uploaded tenant does not match the authenticated tenant."})
        for name, field in fields.items():
            value = row.get(name)
            if field.get("required") and not value:
                errors.append({"row": number, "column": name, "code": "REQUIRED_FIELD_MISSING", "message": "Required value is missing."})
            if value and field.get("options") and value not in field["options"]:
                errors.append({"row": number, "column": name, "value": value, "code": "INVALID_STATUS", "message": "Value is not in the allowed list."})
            if value and (field.get("type") == "email" or field.get("validation_rule") == "EMAIL") and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", value):
                errors.append({"row": number, "column": name, "code": "INVALID_EMAIL", "message": "Invalid email address."})
            if value and field.get("regex_pattern") and not re.fullmatch(field["regex_pattern"], value):
                errors.append({"row": number, "column": name, "code": "REGEX_MISMATCH", "message": "Value does not match the configured validation rule."})
            if value and field.get("min_length") is not None and len(value) < int(field["min_length"]):
                errors.append({"row": number, "column": name, "code": "MIN_LENGTH", "message": "Value is shorter than the configured minimum."})
            if value and field.get("max_length") is not None and len(value) > int(field["max_length"]):
                errors.append({"row": number, "column": name, "code": "MAX_LENGTH", "message": "Value exceeds the configured maximum."})
            if value and field.get("type") == "number":
                try:
                    float(value)
                except (TypeError, ValueError):
                    errors.append({"row": number, "column": name, "code": "INVALID_NUMBER", "message": "Invalid number."})
            if value and field.get("type") == "date":
                try:
                    datetime.fromisoformat(value)
                except ValueError:
                    try:
                        datetime.strptime(value, "%d-%m-%Y")
                    except ValueError:
                        errors.append({"row": number, "column": name, "code": "INVALID_DATE", "message": "Invalid date."})
    return errors


def parse_and_validate(content, filename, cfg, sheet=None, tenant_id=None):
    headers, rows, errors = read_rows(content, filename, sheet)
    mapped, mapping, mapping_errors = map_headers(headers, cfg)
    errors.extend(mapping_errors)
    normalized_rows = []
    for row in rows:
        normalized = {mapping.get(key, key): value for key, value in row.items()}
        normalized_rows.append(normalized)
    errors.extend(validate_rows(mapped, normalized_rows, cfg, tenant_id))
    return mapped, normalized_rows, errors, mapping
