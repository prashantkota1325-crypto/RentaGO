"""Canonical Unique Vendors rate-chart validation and transactional import."""

from datetime import date, datetime, time, timedelta
from io import BytesIO

import openpyxl

from .ids import next_ratecard_id

CANONICAL = (
    ("sr_no", "Sr.No."), ("company_id", "Company Id"), ("legal_name", "Legal Name"),
    ("group_name", "Group"), ("city", "City"), ("state", "State"),
    ("category", "Vehicle Category"), ("vehicle_model", "Vehicle Model"),
    ("package_name", "Package Name"), ("package_rate", "Package Rate"),
    ("pkg_fixed_kms", "Pkg Fixed Km'S"), ("pkg_fixed_hrs", "Pkg Fixed Hr'S"),
    ("extra_km_rate", "Extra Km Rate"), ("extra_hr_rate", "Extra Hr Rate"),
    ("toll_amt", "Toll Amt"), ("parking_amt", "Parking Amt"), ("da", "Da"),
    ("night_allowance_after_10_pm", "Night Allowance After 10 Pm"),
    ("night_allowance_after_11_pm", "Night Allowance After 11 Pm"),
    ("garage_to_garage_kms", "Garage To Garage Km'S"),
    ("garage_to_garage_pct", "Garage To Garage %"),
)
NUMERIC = {key for key, _ in CANONICAL if key not in {"legal_name", "group_name", "city", "state", "category", "vehicle_model", "package_name", "company_id"}}


def normalize_header(value):
    return " ".join(str(value or "").strip().lower().replace("’", "'").split())


def _number(value):
    if value in (None, ""):
        return None
    if isinstance(value, time):
        return value.hour + value.minute / 60
    if isinstance(value, timedelta):
        return value.total_seconds() / 3600
    if isinstance(value, datetime) and value.date() == date(1900, 1, 1):
        return value.hour + value.minute / 60 + value.second / 3600
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise ValueError("must be numeric")
    if result < 0:
        raise ValueError("must not be negative")
    return result


def validate_workbook(content, filename, selected_vendor_id=None, max_rows=10000):
    errors = []
    if not filename.lower().endswith(".xlsx"):
        return [], [{"row": 0, "column": "file", "value": filename, "code": "INVALID_EXTENSION", "message": "Only .xlsx is supported."}]
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        return [], [{"row": 0, "column": "file", "value": "", "code": "INVALID_WORKBOOK", "message": str(exc)}]
    if "Unique Vendors" not in workbook.sheetnames:
        return [], [{"row": 0, "column": "sheet", "value": "", "code": "MISSING_WORKSHEET", "message": "Required worksheet 'Unique Vendors' is missing."}]
    sheet = workbook["Unique Vendors"]
    raw_headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    expected = [normalize_header(label) for _, label in CANONICAL]
    actual = [normalize_header(value) for value in raw_headers]
    if actual[:len(expected)] != expected:
        return [], [{"row": 1, "column": "header", "value": raw_headers, "code": "INVALID_HEADERS", "message": "Unique Vendors headers do not match the canonical format."}]
    rows = []
    seen_serials = set()
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_row=max_rows + 1, values_only=True), 2):
        if not any(value not in (None, "") for value in cells[:len(CANONICAL)]):
            continue
        record = {key: (cells[index] if index < len(cells) else None) for index, (key, _) in enumerate(CANONICAL)}
        for key in ("legal_name", "category", "vehicle_model", "package_name"):
            if not str(record[key] or "").strip():
                errors.append({"row": row_number, "column": key, "value": record[key], "code": "MISSING_REQUIRED", "message": "Required value is missing."})
        for key in NUMERIC:
            try:
                record[key] = _number(record[key])
            except ValueError as exc:
                errors.append({"row": row_number, "column": key, "value": record[key], "code": "INVALID_RATE", "message": str(exc)})
        serial = str(record.get("sr_no") or "").strip()
        if serial and serial in seen_serials:
            errors.append({"row": row_number, "column": "sr_no", "value": serial, "code": "DUPLICATE_SOURCE_ROW", "message": "Duplicate source serial number."})
        if serial:
            seen_serials.add(serial)
        if selected_vendor_id and record.get("company_id") and str(record["company_id"]).strip() != str(selected_vendor_id).strip():
            errors.append({"row": row_number, "column": "company_id", "value": record["company_id"], "code": "VENDOR_MISMATCH", "message": "File Company Id does not match the selected owner."})
        rows.append(record)
    return rows, errors


def import_rows(conn, rows, owner_type, owner_id, tenant_id, vendor_id, source_file):
    cur = conn.cursor()
    created = []
    for row in rows:
        rate_id = next_ratecard_id(cur)
        row_owner_id = row["company_id"] if owner_type == "VENDOR" else owner_id
        row_vendor_id = row["company_id"] if owner_type == "VENDOR" else vendor_id
        cur.execute(
            "INSERT INTO ratecards (rate_card_id,sr_no,company_id,legal_name,group_name,city,state,category,vehicle_model,"
            "package_name,package_rate,pkg_fixed_kms,pkg_fixed_hrs,extra_km_rate,extra_hr_rate,toll_amt,parking_amt,da,"
            "night_allowance_after_10_pm,night_allowance_after_11_pm,garage_to_garage_kms,garage_to_garage_pct,"
            "owner_type,owner_id,tenant_id,vendor_id,status,approval_status,source_file,source_sheet) "
            "VALUES (:1,:2,:3,:4,:5,:6,:7,:8,:9,:10,:11,:12,:13,:14,:15,:16,:17,:18,:19,:20,:21,:22,:23,:24,:25,:26,'DRAFT','PENDING',:27,'Unique Vendors')",
            (rate_id, row["sr_no"], row["company_id"], row["legal_name"], row["group_name"], row["city"], row["state"], row["category"], row["vehicle_model"], row["package_name"], row["package_rate"], row["pkg_fixed_kms"], row["pkg_fixed_hrs"], row["extra_km_rate"], row["extra_hr_rate"], row["toll_amt"], row["parking_amt"], row["da"], row["night_allowance_after_10_pm"], row["night_allowance_after_11_pm"], row["garage_to_garage_kms"], row["garage_to_garage_pct"], owner_type, row_owner_id, tenant_id, row_vendor_id, source_file),
        )
        created.append(rate_id)
    return created
