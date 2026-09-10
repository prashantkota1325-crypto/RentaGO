"""
RentaGO Web - refresh the roles matrix from the workbook.

The first import wrote 'V' for every blank cell of the Roles sheet, losing the
"blank = no access" distinction (SOP 2.2). This script re-reads the Roles
sheet and replaces the roles table rows with correct values:
  F = full, V = view, blank = NULL (no access / sheet hidden).

Usage:
    python scripts/fix_roles.py [workbook_path]
"""

import sys
import os

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from app.config import settings
from app.db import get_connection

ROLE_COLS = ["Super Admin", "CEO", "Sales", "Operations", "Finance", "Vendor Manager",
             "HR", "Customer Service", "Compliance", "Corporate Admin", "Vendor",
             "Driver", "Investor", "HQ"]


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else settings.WORKBOOK_PATH
    print(f"Reading Roles sheet from {path} ...")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "Roles" not in wb.sheetnames:
        print("No Roles sheet found in the workbook.")
        return
    ws = wb["Roles"]

    hr = 5
    hdr = next(ws.iter_rows(min_row=hr, max_row=hr, max_col=30, values_only=True))
    idx = {}
    for c, v in enumerate(hdr, start=1):
        if v is not None and str(v).strip() in ROLE_COLS:
            idx[str(v).strip()] = c
    if not idx:
        print("No role columns found on the Roles sheet header row 5.")
        return

    rows = []
    for row_vals in ws.iter_rows(min_row=hr + 1, max_col=30, values_only=True):
        sheetname = row_vals[0]
        if sheetname is None or str(sheetname).strip() == "":
            continue
        sheetname = str(sheetname).strip()
        for role, col in idx.items():
            val = row_vals[col - 1] if col <= len(row_vals) else None
            if val is None or str(val).strip() == "":
                level = None
            else:
                level = str(val).strip().upper()[:1]
                if level not in ("F", "V"):
                    level = "V"
            rows.append((sheetname, role, level))
    wb.close()

    print(f"Parsed {len(rows)} matrix rows. Replacing roles table ...")
    conn = get_connection()
    cur = conn.cursor()
    cur.execute("DELETE FROM roles")
    for sheet, role, level in rows:
        cur.execute(
            "INSERT INTO roles (sheet, role_code, access_level) VALUES (:1,:2,:3)",
            (sheet, role, level),
        )
    conn.commit()
    cur.execute("SELECT COUNT(*), SUM(CASE WHEN access_level IS NULL THEN 1 ELSE 0 END) "
                "FROM roles")
    total, nulls = cur.fetchone()
    conn.close()
    print(f"Done: {total} rows, {nulls} with NULL (no access).")


if __name__ == "__main__":
    main()
